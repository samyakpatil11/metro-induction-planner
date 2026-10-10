from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path = r"D:\metro\Metro_Induction_Planner_IEEE_Review.docx"
doc = Document(path)

def fmt_run(run, size=9, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rp = run._element.get_or_add_rPr()
    rf = rp.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rp.insert(0, rf)
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")

def replace_text(p, text, size=9.5, bold=False, italic=False):
    for r in list(p.runs):
        p._p.remove(r._r)
    fmt_run(p.add_run(text), size, bold, italic)

def body_paragraph(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(.16)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.widow_control = True
    fmt_run(p.add_run(text), 9.5)
    return p

def subsection(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    fmt_run(p.add_run(text), 9.5, bold=True, italic=True)
    return p

paras = doc.paragraphs
by_text = {p.text.strip(): p for p in paras if p.text.strip()}
# Expand and correct the existing related-work narrative.
replace_text(by_text["Rolling-stock scheduling and maintenance planning have long been studied as coupled optimization problems. Lin and Zhao formulate train assignment together with preventive maintenance while considering operational requirements and depot capacity, illustrating why maintenance availability cannot be optimized independently of service coverage [1]. Yuan et al. integrate a metro timetable, rolling-stock assignment, and short-turning decisions under passenger demand, showing how a richer service model can change fleet allocation choices [2]. More recent work combines timetabling, stop planning, maintenance, and assignment, with explicit circulation and depot constraints [3]."],
"Prior work spans maintenance-feasible fleet rostering and integrated urban rail planning. Zhong et al. generate candidate rolling-stock schedules and then test maintenance feasibility in a second assignment stage; their study uses actual Chinese High-Speed Railway instances and reports operational-cost improvements over manual schedules [1]. Lin and Zhao formulate weekly train assignment with second-level preventive maintenance as a binary optimization model, including maintenance packets and depot capacity [2]. Both studies make maintenance feasibility central, although their time horizons and operating contexts differ from a single-night metro induction decision.")
replace_text(by_text["The reviewed application addresses a narrower horizon and problem. It does not optimize train paths, headways, crew, passenger loads, or multi-day maintenance packets. Instead, it assigns each train one nightly state and covers fixed line-level demand. This narrower model is appropriate for a prototype focused on induction decisions, but the scope must remain explicit: results cannot be interpreted as a complete metro operating plan."],
"Metro studies integrate additional operational decisions. Yuan et al. jointly optimize a metro timetable, rolling-stock assignment, and short-turning under passenger demand [3]. Yin et al. extend the problem to connected urban lines and depots using a space-time network, a bi-objective mixed-integer model, and Benders decomposition; their evaluation uses Beijing rail-network instances [4]. Li et al. combine timetabling, stop planning, rolling-stock maintenance, and assignment, using an adaptive large-neighborhood search for a larger integrated problem [5]. These models represent a broader operational scope than the reviewed planner.")
replace_text(by_text["The solver uses Google OR-Tools CP-SAT, a constraint-programming and satisfiability-based optimization engine whose documented scheduling examples encode assignment variables and logical constraints [4]. The choice is consistent with a problem containing discrete train states, hard safety exclusions, coverage requirements, and weighted preferences."],
"Predictive maintenance research adds a complementary data perspective. Davari et al. survey data-driven predictive-maintenance methods in the railway industry and discuss the role of data, task definition, and evaluation metrics [6]. The reviewed planner uses a logistic-regression risk score, but its bundled fault histories are synthetic; therefore, its risk estimates should be treated as a demonstration capability rather than a validated predictive result. The assignment solver uses Google OR-Tools CP-SAT, which supports discrete variables and logical scheduling constraints [7].")

# Correct section numbering for inserted comparative synthesis.
renumber = {
    "III. REVIEW METHOD":"IV. REVIEW METHOD",
    "IV. SYSTEM ARCHITECTURE":"V. SYSTEM ARCHITECTURE",
    "V. ASSIGNMENT MODEL":"VI. ASSIGNMENT MODEL",
    "VI. RESULTS AND TECHNICAL ASSESSMENT":"VII. RESULTS AND TECHNICAL ASSESSMENT",
    "VII. RECOMMENDATIONS":"VIII. RECOMMENDATIONS",
    "VIII. CONCLUSION":"IX. CONCLUSION",
}
for old, new in renumber.items():
    if old in by_text:
        replace_text(by_text[old], new, 9.5, True)

# Add comparative table and place it directly after the related-work paragraphs.
caption = doc.add_paragraph()
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_before = Pt(3)
caption.paragraph_format.space_after = Pt(2)
caption.paragraph_format.keep_with_next = True
fmt_run(caption.add_run("TABLE I  COMPARATIVE REVIEW OF RELEVANT STUDIES"), 7.5, True)

table = doc.add_table(rows=1, cols=3)
table.autofit = False
table.style = "Table Grid"
widths = [Inches(.80), Inches(1.23), Inches(1.37)]
headers = ["Study", "Scope and method", "Comparison with planner"]
for cell, text, width in zip(table.rows[0].cells, headers, widths):
    cell.width = width
    cell.text = text
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs: fmt_run(r, 7.2, True)
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "D9E2F3"); tcPr.append(shd)
rows = [
    ("Zhong et al. [1]", "Maintenance-feasible HSR rostering; two-stage MILP/assignment; actual cases.", "Multi-stage maintenance and real-data evidence; broader timetable context."),
    ("Lin and Zhao [2]", "Weekly train assignment and preventive maintenance; binary optimization.", "Multi-day maintenance packets exceed the planner's nightly horizon."),
    ("Yuan et al. [3]", "Metro timetable, fleet assignment, short-turning; demand-aware integrated model.", "Adds passenger demand and timetable decisions absent from the planner."),
    ("Yin et al. [4]", "Multi-line urban rail allocation and timetabling; bi-objective MILP/Benders.", "Models connected depots and network-scale allocation."),
    ("Li et al. [5]", "Timetable, stop plan, maintenance, and fleet assignment; MILP/ALNS.", "Broader integrated scope, with more operating inputs and complexity."),
    ("Davari et al. [6]", "Survey of data-driven railway predictive-maintenance methods and evaluation.", "Highlights data and validation needs for the planner's risk model."),
    ("Reviewed system", "Nightly service/standby/maintenance assignment; CP-SAT, fixed demand, risk score.", "Operator-facing safety checks; synthetic risk history and no field validation."),
]
for rowdata in rows:
    cells = table.add_row().cells
    for cell, text, width in zip(cells, rowdata, widths):
        cell.width = width
        cell.text = text
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = .95
            for r in p.runs: fmt_run(r, 7.0)
    trPr = table.rows[-1]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))

# Place caption + comparison table following existing literature paragraphs.
anchor = next(p for p in doc.paragraphs if p.text.startswith("Predictive maintenance research adds"))
anchor._p.addnext(caption._p)
caption._p.addnext(table._tbl)

# New explicit synthesis and gap section before methods.
method_heading = next(p for p in doc.paragraphs if p.text == "IV. REVIEW METHOD")
new_items = [
    ("heading", "III. COMPARATIVE SYNTHESIS AND RESEARCH GAPS"),
    ("body", "The literature shows a progression from maintenance-aware rostering [1], [2] toward integrated metro planning that couples vehicle allocation with timetables, passenger demand, line interactions, and depot operations [3]–[5]. The reviewed planner occupies a deliberately narrower position: it translates a compact nightly assignment model into an explainable, approval-based interface. Its contribution is therefore system integration and operator workflow, rather than a new optimization formulation or a claim of superior objective values."),
    ("sub", "A. Scope and time-horizon gap"),
    ("body", "The planner assigns one state per train for one night and targets fixed line coverage. It omits timetable and headway feasibility, passenger demand, crew, rolling-stock circulation across trips, multi-day maintenance packets, uncertain failures, and network disruption recovery. Integrated models in [3]–[5] demonstrate that these decisions interact. Future work should either justify the nightly abstraction against depot practice or extend the model incrementally, with independent feasibility checks for each added operating constraint."),
    ("sub", "B. Empirical and predictive-validity gap"),
    ("body", "The current logistic-regression model is trained on generated demo histories whose fault process is encoded by a synthetic rule. The self-check confirms that the model trains and produces nonconstant bounded scores, but does not establish predictive accuracy or calibration. Railway predictive-maintenance research emphasizes task-specific data and evaluation [6]. A real-data study should use temporally separated train-level validation, report class-sensitive metrics and calibration, compare simple baselines, quantify uncertainty, and test drift. Until then, risk scores must remain advisory and must not override deterministic safety exclusions."),
    ("sub", "C. Operational transfer and human-factors gap"),
    ("body", "Published scheduling studies evaluate actual or network-scale instances [1], [4], whereas the reviewed app has no reported pilot with depot officers or maintainers. Configuration values for service demand, maintenance capacity, overhaul mileage, certificate windows, cleaning, and shunting need authority approval. Scenario tests should include degraded fleet capacity, stale records, conflicting officer pins, solver timeouts, and faults reported after plan creation. Usability evaluation should assess whether explanations, warnings, plan status, and bulk-assignment previews support correct decisions under shift pressure."),
    ("sub", "D. Governance and deployment gap"),
    ("body", "The code contains demo users and a fallback token-signing secret, and the local deployment uses SQLite. Before production, the operator needs controlled identities, secret rotation, backup and recovery, access and decision audit trails, versioned rule sets, data provenance, and a documented incident process. A shadow-mode trial should compare recommendations against expert decisions without controlling dispatch, then define acceptance thresholds and rollback criteria."),
    ("body", "These gaps delimit the claims of this review: automated checks establish software invariants on supplied fixtures, while operational performance, model validity, and organizational fitness remain open research and engineering questions."),
]
for kind, text in new_items:
    p = doc.add_paragraph()
    if kind == "heading":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        fmt_run(p.add_run(text), 9.5, True)
    elif kind == "sub":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        fmt_run(p.add_run(text), 9.5, True, True)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(.16)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.widow_control = True
        fmt_run(p.add_run(text), 9.5)
    method_heading._p.addprevious(p._p)

# Replace bibliography with directly cited technical literature and solver documentation.
ref_head = next(p for p in doc.paragraphs if p.text.strip().upper() == "REFERENCES")
for p in list(doc.paragraphs):
    if p._p is ref_head._p:
        after = True
        p._p.getparent().remove(p._p)
        continue
    if 'after' in locals() and after:
        p._p.getparent().remove(p._p)
refs = [
"[1] Q. Zhong, R. M. Lusby, J. Larsen, Y. Zhang, and Q. Peng, “Rolling stock scheduling with maintenance requirements at the Chinese High-Speed Railway,” Transportation Research Part B: Methodological, vol. 126, pp. 24–44, Aug. 2019, doi: 10.1016/j.trb.2019.05.013.",
"[2] B. Lin and Y. Zhao, “Synchronized optimization of EMU train assignment and second-level preventive maintenance scheduling,” Reliability Engineering & System Safety, vol. 215, Art. no. 107893, 2021, doi: 10.1016/j.ress.2021.107893.",
"[3] J. Yuan, Y. Gao, S. Li, P. Liu, and L. Yang, “Integrated optimization of train timetable, rolling stock assignment and short-turning strategy for a metro line,” European Journal of Operational Research, vol. 301, no. 3, pp. 855–874, 2022, doi: 10.1016/j.ejor.2021.11.019.",
"[4] J. Yin, F. Pu, L. Yang, A. D’Ariano, and Z. Wang, “Integrated optimization of rolling stock allocation and train timetables for urban rail transit networks: A Benders decomposition approach,” Transportation Research Part B: Methodological, vol. 176, Art. no. 102815, 2023, doi: 10.1016/j.trb.2023.102815.",
"[5] T. Li, D. Li, H. Ji, and J. Ma, “Integrated optimization of train timetabling, stop planning, rolling stock maintenance and assignment with flexible train composition and short-turning strategies,” Transportation Research Part E: Logistics and Transportation Review, vol. 207, Art. no. 104566, 2026, doi: 10.1016/j.tre.2025.104566.",
"[6] N. Davari, B. Veloso, G. d. A. Costa, P. M. Pereira, R. P. Ribeiro, and J. Gama, “A survey on data-driven predictive maintenance for the railway industry,” Sensors, vol. 21, no. 17, Art. no. 5739, 2021, doi: 10.3390/s21175739.",
"[7] Google, “Employee scheduling | OR-Tools,” Google for Developers. [Online]. Available: https://developers.google.com/optimization/scheduling/employee_scheduling. Accessed: Oct. 4, 2026."
]
h = doc.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
h.paragraph_format.space_before = Pt(7)
h.paragraph_format.space_after = Pt(4)
h.paragraph_format.keep_with_next = True
fmt_run(h.add_run("REFERENCES"), 9.5, True)
for text in refs:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(.18)
    p.paragraph_format.first_line_indent = Inches(-.18)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    fmt_run(p.add_run(text), 8)

doc.save(path)
print(path)

